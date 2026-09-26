import tkinter as tk
from tkinter import ttk, messagebox, simpledialog
import openpyxl
import os
import re

# Excel 文件路径（默认保存在桌面）
EXCEL_PATH = os.path.join(os.path.expanduser("~"), "Desktop", "数据收集表.xlsx")

def init_excel():
    """初始化Excel文件，创建表头（若文件不存在）"""
    if not os.path.exists(EXCEL_PATH):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "数据收集"
        # 设置表头
        headers = ["序号", "姓名", "电话", "地址", "备注", "录入时间"]
        ws.append(headers)
        # 设置表头样式（加粗）
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)
        wb.save(EXCEL_PATH)
    return openpyxl.load_workbook(EXCEL_PATH)

def validate_phone(phone):
    """校验手机号格式（11位数字）"""
    pattern = r"^1[3-9]\d{9}$"
    return re.match(pattern, phone) is not None

def add_data_to_excel(name, phone, address, remark):
    """将数据追加到Excel"""
    try:
        wb = init_excel()
        ws = wb.active
        
        # 获取当前最大行号（序号）
        max_row = ws.max_row
        serial_num = max_row  # 序号从1开始，表头占1行
        
        # 获取当前时间（简化版，也可导入datetime模块）
        import datetime
        now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        
        # 组装数据行
        row_data = [serial_num, name, phone, address, remark, now]
        ws.append(row_data)
        
        # 保存文件
        wb.save(EXCEL_PATH)
        wb.close()
        return True
    except Exception as e:
        messagebox.showerror("错误", f"数据保存失败：{str(e)}")
        return False

def submit_data():
    """提交按钮触发的逻辑"""
    # 获取输入框内容
    name = entry_name.get().strip()
    phone = entry_phone.get().strip()
    address = entry_address.get("1.0", tk.END).strip()  # 多行文本框
    remark = entry_remark.get().strip()
    
    # 必填项校验
    if not name:
        messagebox.warning("提示", "姓名不能为空！")
        return
    if not phone:
        messagebox.warning("提示", "电话不能为空！")
        return
    
    # 手机号格式校验
    if not validate_phone(phone):
        messagebox.warning("提示", "手机号格式错误（请输入11位有效手机号）！")
        return
    
    # 写入Excel
    if add_data_to_excel(name, phone, address, remark):
        messagebox.showinfo("成功", "数据已保存到Excel！")
        # 清空输入框
        entry_name.delete(0, tk.END)
        entry_phone.delete(0, tk.END)
        entry_address.delete("1.0", tk.END)
        entry_remark.delete(0, tk.END)

def view_data():
    """查看已收集的数据"""
    try:
        wb = init_excel()
        ws = wb.active
        max_row = ws.max_row
        
        if max_row == 1:
            messagebox.showinfo("提示", "暂无数据！")
            return
        
        # 新建窗口展示数据
        view_window = tk.Toplevel(root)
        view_window.title("已收集数据")
        view_window.geometry("800x400")
        
        # 创建表格
        tree = ttk.Treeview(view_window, columns=("序号", "姓名", "电话", "地址", "备注", "录入时间"), show="headings")
        tree.heading("序号", text="序号")
        tree.heading("姓名", text="姓名")
        tree.heading("电话", text="电话")
        tree.heading("地址", text="地址")
        tree.heading("备注", text="备注")
        tree.heading("录入时间", text="录入时间")
        
        # 设置列宽
        tree.column("序号", width=50)
        tree.column("姓名", width=100)
        tree.column("电话", width=120)
        tree.column("地址", width=200)
        tree.column("备注", width=150)
        tree.column("录入时间", width=180)
        
        # 填充数据
        for row in range(2, max_row + 1):
            row_data = [ws.cell(row=row, column=col).value for col in range(1, 7)]
            tree.insert("", tk.END, values=row_data)
        
        # 滚动条
        scrollbar = ttk.Scrollbar(view_window, orient=tk.VERTICAL, command=tree.yview)
        tree.configure(yscrollcommand=scrollbar.set)
        tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        wb.close()
    except Exception as e:
        messagebox.showerror("错误", f"读取数据失败：{str(e)}")

def clear_excel():
    """清空Excel数据（保留表头）"""
    if messagebox.askyesno("确认", "是否清空所有收集的数据？（表头将保留）"):
        try:
            wb = init_excel()
            ws = wb.active
            # 删除除表头外的所有行
            max_row = ws.max_row
            if max_row > 1:
                ws.delete_rows(2, max_row - 1)
            wb.save(EXCEL_PATH)
            wb.close()
            messagebox.showinfo("成功", "数据已清空！")
        except Exception as e:
            messagebox.showerror("错误", f"清空数据失败：{str(e)}")

# 主窗口配置
if __name__ == "__main__":
    root = tk.Tk()
    root.title("数据收集桌面软件")
    root.geometry("600x450")
    root.resizable(False, False)  # 禁止窗口缩放

    # 样式配置
    ttk.Style().configure("TLabel", font=("微软雅黑", 10))
    ttk.Style().configure("TButton", font=("微软雅黑", 10))
    ttk.Style().configure("TEntry", font=("微软雅黑", 10))

    # 1. 姓名输入区
    label_name = ttk.Label(root, text="姓名（必填）：")
    label_name.place(x=50, y=30, width=100, height=30)
    entry_name = ttk.Entry(root)
    entry_name.place(x=160, y=30, width=350, height=30)

    # 2. 电话输入区
    label_phone = ttk.Label(root, text="电话（必填）：")
    label_phone.place(x=50, y=80, width=100, height=30)
    entry_phone = ttk.Entry(root)
    entry_phone.place(x=160, y=80, width=350, height=30)

    # 3. 地址输入区（多行文本）
    label_address = ttk.Label(root, text="地址：")
    label_address.place(x=50, y=130, width=100, height=30)
    entry_address = tk.Text(root, font=("微软雅黑", 10))
    entry_address.place(x=160, y=130, width=350, height=100)

    # 4. 备注输入区
    label_remark = ttk.Label(root, text="备注：")
    label_remark.place(x=50, y=250, width=100, height=30)
    entry_remark = ttk.Entry(root)
    entry_remark.place(x=160, y=250, width=350, height=30)

    # 5. 功能按钮区
    btn_submit = ttk.Button(root, text="提交数据", command=submit_data)
    btn_submit.place(x=80, y=310, width=120, height=40)

    btn_view = ttk.Button(root, text="查看数据", command=view_data)
    btn_view.place(x=230, y=310, width=120, height=40)

    btn_clear = ttk.Button(root, text="清空数据", command=clear_excel)
    btn_clear.place(x=380, y=310, width=120, height=40)

    # 初始化Excel
    init_excel()

    # 运行主循环
    root.mainloop()
