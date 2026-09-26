from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import os
import sys

# 解决EXE运行时的路径问题
def get_resource_path(relative_path):
    if hasattr(sys, '_MEIPASS'):
        return os.path.join(sys._MEIPASS, relative_path)
    return os.path.join(os.path.abspath("."), relative_path)

def get_contact_info():
    """获取用户输入的姓名、电话、地址"""
    print("===== 联系人信息录入 =====")
    while True:
        name = input("请输入姓名：").strip()
        if name:
            break
        print("姓名不能为空，请重新输入！")
    
    while True:
        phone = input("请输入电话：").strip()
        if phone:
            break
        print("电话不能为空，请重新输入！")
    
    address = input("请输入地址（可选）：").strip()
    return {"姓名": name, "电话": phone, "地址": address if address else "未填写"}

def create_word_file(contact_info):
    """生成Word文档（保存到桌面）"""
    # 固定保存到桌面，避免EXE路径混乱
    desktop_path = os.path.join(os.path.expanduser("~"), "Desktop")
    save_path = os.path.join(desktop_path, "联系人信息.docx")
    
    doc = Document()
    # 设置标题
    title = doc.add_heading("联系人信息", level=1)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    # 添加信息
    for key, value in contact_info.items():
        para = doc.add_paragraph()
        run = para.add_run(f"{key}：{value}")
        run.font.name = "微软雅黑"
        run.font.size = Pt(12)
        para.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    
    # 保存文件
    try:
        doc.save(save_path)
        if os.path.exists(save_path):
            print(f"\n✅ Word文档已生成！")
            print(f"📁 文件路径：{save_path}")
        else:
            print("❌ 文档生成失败！")
    except Exception as e:
        print(f"❌ 保存失败：{str(e)}")

if __name__ == "__main__":
    contact = get_contact_info()
    create_word_file(contact)
    input("\n按回车键退出程序...")
